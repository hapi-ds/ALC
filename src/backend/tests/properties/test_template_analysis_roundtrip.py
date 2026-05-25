"""Property-based tests for template analysis round-trip.

Property 1: Template Analysis Round-Trip
For any valid .docx file with headings and placeholder markers, the
analyze_template() method correctly extracts the structure and
detect_placeholders() finds all embedded markers.

Verifies:
- All headings inserted into a .docx appear in section_hierarchy
- All placeholder markers embedded in the document are detected
- The structural analysis is consistent with the input document

**Validates: Requirements 1.2, 1.3, 10.1**

References:
    - Design: .kiro/specs/Step_5-4_ai-document-generator-template-based/design.md
    - Requirements: .kiro/specs/Step_5-4_ai-document-generator-template-based/requirements.md
"""

import io
from unittest.mock import AsyncMock, MagicMock

import hypothesis.strategies as st
from docx import Document as DocxDocument
from hypothesis import given, settings

from alcoabase.services.template_analysis import (
    TemplateAnalysisService,
)


# ---------------------------------------------------------------------------
# Hypothesis Strategies
# ---------------------------------------------------------------------------

# Valid placeholder identifiers: 1-50 uppercase letters and underscores
PLACEHOLDER_IDENTIFIERS = st.from_regex(r"[A-Z_]{1,50}", fullmatch=True)

# Valid placeholder parameters: 1-100 printable characters (no closing braces,
# no control characters that are invalid in XML)
PLACEHOLDER_PARAMETERS = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S", "Z"),
        min_codepoint=32,
        max_codepoint=126,
        blacklist_characters="}",
    ),
    min_size=1,
    max_size=100,
).filter(lambda s: len(s.strip()) > 0)

# Heading levels supported by the service (1-4)
HEADING_LEVELS = st.integers(min_value=1, max_value=4)

# Heading text: non-empty printable ASCII strings (no newlines)
HEADING_TEXT = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S"),
        min_codepoint=32,
        max_codepoint=126,
    ),
    min_size=1,
    max_size=80,
).filter(lambda s: len(s.strip()) > 0)

# Body paragraph text (non-empty, no placeholder-like patterns)
BODY_TEXT = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S"),
        min_codepoint=32,
        max_codepoint=126,
    ),
    min_size=1,
    max_size=200,
).filter(lambda s: "{{" not in s and len(s.strip()) > 0)


@st.composite
def placeholder_marker(draw: st.DrawFn) -> str:
    """Generate a valid placeholder marker string."""
    identifier = draw(PLACEHOLDER_IDENTIFIERS)
    has_param = draw(st.booleans())
    if has_param:
        parameter = draw(PLACEHOLDER_PARAMETERS)
        return f"{{{{{identifier}:{parameter}}}}}"
    return f"{{{{{identifier}}}}}"


@st.composite
def document_structure(draw: st.DrawFn) -> dict:
    """Generate a document structure with headings and placeholders.

    Returns a dict with:
        headings: list of (level, text) tuples
        placeholders: list of placeholder marker strings
        body_paragraphs: list of body text strings (may contain placeholders)
    """
    num_headings = draw(st.integers(min_value=1, max_value=8))
    headings = [
        (draw(HEADING_LEVELS), draw(HEADING_TEXT))
        for _ in range(num_headings)
    ]

    num_placeholders = draw(st.integers(min_value=0, max_value=5))
    placeholders = [draw(placeholder_marker()) for _ in range(num_placeholders)]

    # Generate body paragraphs, some of which contain placeholders
    body_paragraphs = []
    for ph in placeholders:
        prefix = draw(BODY_TEXT)
        body_paragraphs.append(f"{prefix} {ph}")

    return {
        "headings": headings,
        "placeholders": placeholders,
        "body_paragraphs": body_paragraphs,
    }


# ---------------------------------------------------------------------------
# Helper: Build a .docx file from a document structure
# ---------------------------------------------------------------------------


def build_docx_bytes(structure: dict) -> bytes:
    """Create a .docx file in memory from the given structure.

    Args:
        structure: Dict with headings, placeholders, and body_paragraphs.

    Returns:
        Raw bytes of the generated .docx file.
    """
    doc = DocxDocument()

    headings = structure["headings"]
    body_paragraphs = structure["body_paragraphs"]

    # Distribute body paragraphs across sections
    body_idx = 0
    for level, text in headings:
        doc.add_heading(text, level=level)
        # Add a body paragraph after each heading if available
        if body_idx < len(body_paragraphs):
            doc.add_paragraph(body_paragraphs[body_idx])
            body_idx += 1

    # Add remaining body paragraphs after the last heading
    while body_idx < len(body_paragraphs):
        doc.add_paragraph(body_paragraphs[body_idx])
        body_idx += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Helper: Create TemplateAnalysisService instance (no DB/storage needed)
# ---------------------------------------------------------------------------


def _make_service() -> TemplateAnalysisService:
    """Create a TemplateAnalysisService with mocked dependencies.

    Only analyze_template() and detect_placeholders() are pure methods
    that don't require real DB or storage connections.
    """
    session_factory = MagicMock()
    storage_service = AsyncMock()
    job_tracker = AsyncMock()
    return TemplateAnalysisService(
        session_factory=session_factory,
        storage_service=storage_service,
        job_tracker=job_tracker,
    )


# ---------------------------------------------------------------------------
# Property 1: All headings appear in section_hierarchy
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(structure=document_structure())
def test_all_headings_appear_in_section_hierarchy(
    structure: dict,
) -> None:
    """For any valid .docx with headings, analyze_template() SHALL extract
    all headings into section_hierarchy with correct heading text and levels.

    **Validates: Requirements 1.2**
    """
    service = _make_service()
    docx_bytes = build_docx_bytes(structure)
    analysis = service.analyze_template(docx_bytes)

    expected_headings = structure["headings"]

    # All headings should appear in section_hierarchy
    assert len(analysis.section_hierarchy) == len(expected_headings)

    for i, (expected_level, expected_text) in enumerate(expected_headings):
        section = analysis.section_hierarchy[i]
        assert section.heading == expected_text
        assert section.level == expected_level
        assert section.position == i


# ---------------------------------------------------------------------------
# Property 1: All placeholder markers are detected
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(structure=document_structure())
def test_all_placeholder_markers_detected(
    structure: dict,
) -> None:
    """For any valid .docx with placeholder markers, analyze_template() SHALL
    detect all embedded markers via detect_placeholders().

    **Validates: Requirements 1.3, 10.1**
    """
    service = _make_service()
    docx_bytes = build_docx_bytes(structure)
    analysis = service.analyze_template(docx_bytes)

    expected_placeholders = set(structure["placeholders"])

    if not expected_placeholders:
        # No placeholders expected — analysis should have empty or no markers
        # (some may still be detected if heading text accidentally matches)
        return

    # All expected placeholders should be found in the analysis
    detected_markers = {p["marker"] for p in analysis.placeholder_markers}
    for expected in expected_placeholders:
        assert expected in detected_markers, (
            f"Expected placeholder {expected!r} not found in detected markers. "
            f"Detected: {detected_markers}"
        )


# ---------------------------------------------------------------------------
# Property 1: detect_placeholders() finds all markers in text
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    identifiers=st.lists(PLACEHOLDER_IDENTIFIERS, min_size=1, max_size=5),
    has_params=st.lists(st.booleans(), min_size=1, max_size=5),
)
def test_detect_placeholders_finds_all_markers_in_text(
    identifiers: list[str],
    has_params: list[bool],
) -> None:
    """detect_placeholders() SHALL find all {{IDENTIFIER}} and
    {{IDENTIFIER:parameter}} patterns embedded in arbitrary text.

    **Validates: Requirements 1.3, 10.1**
    """
    service = _make_service()
    # Build text with known placeholders
    markers = []
    text_parts = ["Some prefix text. "]

    for i, identifier in enumerate(identifiers):
        use_param = has_params[i] if i < len(has_params) else False
        if use_param:
            marker = f"{{{{{identifier}:test_param}}}}"
        else:
            marker = f"{{{{{identifier}}}}}"
        markers.append(marker)
        text_parts.append(f"Content before {marker} content after. ")

    text = "".join(text_parts)

    results = service.detect_placeholders(text)
    detected_markers = [r["marker"] for r in results]

    # Every marker we inserted should be detected
    for marker in markers:
        assert marker in detected_markers, (
            f"Marker {marker!r} not found in results. "
            f"Detected: {detected_markers}"
        )


# ---------------------------------------------------------------------------
# Property 1: detect_placeholders() returns correct identifier and parameter
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    identifier=PLACEHOLDER_IDENTIFIERS,
    parameter=st.one_of(st.none(), PLACEHOLDER_PARAMETERS),
)
def test_detect_placeholders_extracts_identifier_and_parameter(
    identifier: str,
    parameter: str | None,
) -> None:
    """detect_placeholders() SHALL correctly extract the identifier and
    optional parameter from each detected placeholder marker.

    **Validates: Requirements 1.3**
    """
    service = _make_service()
    if parameter is not None:
        marker = f"{{{{{identifier}:{parameter}}}}}"
    else:
        marker = f"{{{{{identifier}}}}}"

    text = f"Before {marker} after"
    results = service.detect_placeholders(text)

    assert len(results) >= 1
    # Find our specific marker in results
    matching = [r for r in results if r["marker"] == marker]
    assert len(matching) == 1, f"Expected exactly one match for {marker!r}"

    result = matching[0]
    assert result["identifier"] == identifier
    if parameter is not None:
        assert result["parameter"] == parameter
    else:
        assert result["parameter"] == ""


# ---------------------------------------------------------------------------
# Property 1: section_hierarchy total_sections matches count
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(structure=document_structure())
def test_total_sections_matches_hierarchy_length(
    structure: dict,
) -> None:
    """analyze_template() SHALL report total_sections equal to the length
    of section_hierarchy.

    **Validates: Requirements 1.2**
    """
    service = _make_service()
    docx_bytes = build_docx_bytes(structure)
    analysis = service.analyze_template(docx_bytes)

    assert analysis.total_sections == len(analysis.section_hierarchy)
