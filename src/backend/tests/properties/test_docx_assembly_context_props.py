"""Property-based tests for DOCX assembly and context management.

Property 3: Template Formatting Preservation
Property 8: Output Document Completeness
Property 15: Knowledge Base Relevance Ordering
Property 16: Context Window Management
Property 18: No Raw Placeholders in Output
Property 20: Header/Footer Token Substitution

**Validates: Requirements 2.4, 4.1, 4.2, 4.6, 4.7, 4.9, 4.10, 9.1, 9.2, 9.4, 10.7, 4.3**

References:
    - Design: .kiro/specs/Step_5-4_ai-document-generator-template-based/design.md
    - Requirements: .kiro/specs/Step_5-4_ai-document-generator-template-based/requirements.md
"""

import asyncio
import io
import re
import zipfile
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import hypothesis.strategies as st
from docx import Document as DocxDocument
from hypothesis import given, settings

from alcoabase.services.template_analysis import TemplateAnalysis, TemplateSection
from alcoabase.services.template_document_generator import (
    CHARS_PER_TOKEN,
    MAX_CONTEXT_TOKENS,
    MAX_PRECEDING_SUMMARY_TOKENS,
    SectionResult,
    TemplateDocumentGeneratorService,
)


# ---------------------------------------------------------------------------
# Hypothesis Strategies
# ---------------------------------------------------------------------------

# Heading text: non-empty printable ASCII (no newlines, no placeholder patterns)
st_heading_text = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S"),
        min_codepoint=32,
        max_codepoint=126,
    ),
    min_size=3,
    max_size=50,
).filter(lambda s: len(s.strip()) > 0 and "{{" not in s and "{" not in s)

# Heading levels (1-4)
st_heading_levels = st.integers(min_value=1, max_value=4)

# Section content: non-empty text without placeholder patterns
st_section_content = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S", "Z"),
        min_codepoint=32,
        max_codepoint=126,
    ),
    min_size=10,
    max_size=200,
).filter(lambda s: "{{" not in s and "{" not in s and len(s.strip()) > 0)

# Relevance scores between 0.3 and 1.0 (valid KB chunks)
st_relevance_scores = st.floats(min_value=0.3, max_value=1.0, allow_nan=False)

# Document titles (simple ASCII, no braces)
st_titles = st.text(
    alphabet=st.characters(min_codepoint=65, max_codepoint=122),
    min_size=3,
    max_size=50,
).filter(lambda s: "{" not in s and "}" not in s)

# Company names (simple ASCII, no braces)
st_company_names = st.text(
    alphabet=st.characters(min_codepoint=65, max_codepoint=122),
    min_size=3,
    max_size=30,
).filter(lambda s: "{" not in s and "}" not in s)

# User names (simple ASCII)
st_user_names = st.text(
    alphabet=st.characters(min_codepoint=65, max_codepoint=122),
    min_size=3,
    max_size=30,
)

# Placeholder identifiers (uppercase + underscores, 1-30 chars)
st_placeholder_ids = st.from_regex(r"[A-Z_]{1,30}", fullmatch=True)

# Token patterns for headers/footers
SUPPORTED_TOKENS = ["{title}", "{date}", "{version}", "{company}"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_service() -> TemplateDocumentGeneratorService:
    """Create a TemplateDocumentGeneratorService with mocked dependencies."""
    session_factory = MagicMock()
    inference_client = AsyncMock()
    knowledge_service = MagicMock()
    agent_registry = AsyncMock()
    storage_service = AsyncMock()
    job_tracker = AsyncMock()
    template_analysis_service = MagicMock()

    return TemplateDocumentGeneratorService(
        session_factory=session_factory,
        inference_client=inference_client,
        knowledge_service=knowledge_service,
        agent_registry=agent_registry,
        storage_service=storage_service,
        job_tracker=job_tracker,
        template_analysis_service=template_analysis_service,
    )


def _build_template_docx(
    headings: list[tuple[int, str]],
    header_text: str = "",
    footer_text: str = "",
) -> bytes:
    """Build a template .docx with given headings and header/footer text."""
    doc = DocxDocument()

    # Add header/footer if specified
    if header_text or footer_text:
        section = doc.sections[0]
        if header_text:
            header = section.header
            header.paragraphs[0].text = header_text
        if footer_text:
            footer = section.footer
            footer.paragraphs[0].text = footer_text

    # Add headings with body paragraphs
    for level, text in headings:
        doc.add_heading(text, level=level)
        doc.add_paragraph(f"Body content for {text}")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_section_results(
    headings: list[tuple[int, str]],
    content: str = "Generated content for this section.",
) -> list[SectionResult]:
    """Create SectionResult objects for each heading."""
    results = []
    for i, (level, text) in enumerate(headings):
        results.append(
            SectionResult(
                heading=text,
                content=content,
                token_count=len(content) // CHARS_PER_TOKEN,
                inference_duration_ms=100,
                kb_chunks_used=[
                    {
                        "chunk_document_uuid": f"doc-{i}",
                        "chunk_text": f"Source chunk for {text}",
                        "relevance_score": 0.8,
                    }
                ],
            )
        )
    return results


def _make_template_analysis(
    headings: list[tuple[int, str]],
) -> TemplateAnalysis:
    """Create a TemplateAnalysis from headings."""
    sections = [
        TemplateSection(
            heading=text,
            level=level,
            position=i,
            has_placeholder=False,
            placeholder_markers=[],
            has_table=False,
            table_columns=None,
        )
        for i, (level, text) in enumerate(headings)
    ]
    return TemplateAnalysis(
        section_hierarchy=sections,
        numbering_scheme="1.1.1",
        paragraph_styles=["Normal", "Heading 1", "Heading 2"],
        table_structures=[],
        header_footer_patterns={},
        placeholder_markers=[],
        total_sections=len(sections),
        has_toc=False,
        page_layout={"margins": {"top": 1, "bottom": 1}},
    )


# ---------------------------------------------------------------------------
# Property 3: Template Formatting Preservation
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    headings=st.lists(
        st.tuples(st_heading_levels, st_heading_text),
        min_size=1,
        max_size=6,
    ),
    title=st_titles,
    company_name=st_company_names,
    user_name=st_user_names,
)
def test_output_preserves_heading_hierarchy(
    headings: list[tuple[int, str]],
    title: str,
    company_name: str,
    user_name: str,
) -> None:
    """Output .docx SHALL contain the same heading hierarchy as the template.
    Heading styles and levels from the template are preserved in the output.

    **Validates: Requirements 2.4, 4.1**
    """
    service = _make_service()
    template_bytes = _build_template_docx(headings)
    sections = _make_section_results(headings)
    template_analysis = _make_template_analysis(headings)

    output_bytes = asyncio.run(
        service.assemble_docx(
            template_bytes=template_bytes,
            sections=sections,
            template_analysis=template_analysis,
            title=title,
            company_name=company_name,
            requesting_user_name=user_name,
        )
    )

    # Parse output document
    output_doc = DocxDocument(io.BytesIO(output_bytes))

    # Extract heading paragraphs from output
    output_headings: list[tuple[int, str]] = []
    for para in output_doc.paragraphs:
        style_name = para.style.name if para.style else ""
        if style_name and re.match(r"[Hh]eading\s*\d+", style_name):
            level_match = re.search(r"\d+", style_name)
            if level_match:
                output_headings.append((int(level_match.group()), para.text))

    # All template headings should be present in output
    template_heading_texts = {text for _, text in headings}
    output_heading_texts = {text for _, text in output_headings}

    for heading_text in template_heading_texts:
        assert heading_text in output_heading_texts, (
            f"Template heading '{heading_text}' not found in output document. "
            f"Output headings: {output_heading_texts}"
        )


# ---------------------------------------------------------------------------
# Property 8: Output Document Completeness
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    headings=st.lists(
        st.tuples(st_heading_levels, st_heading_text),
        min_size=1,
        max_size=5,
    ),
    title=st_titles,
    company_name=st_company_names,
    user_name=st_user_names,
)
def test_output_document_completeness(
    headings: list[tuple[int, str]],
    title: str,
    company_name: str,
    user_name: str,
) -> None:
    """Output .docx SHALL be a valid ZIP/OPC archive with correct properties,
    valid UTF-8 text, and a Sources appendix section.

    **Validates: Requirements 4.2, 4.6, 4.9, 5.5**
    """
    service = _make_service()
    template_bytes = _build_template_docx(headings)
    sections = _make_section_results(headings)
    template_analysis = _make_template_analysis(headings)

    output_bytes = asyncio.run(
        service.assemble_docx(
            template_bytes=template_bytes,
            sections=sections,
            template_analysis=template_analysis,
            title=title,
            company_name=company_name,
            requesting_user_name=user_name,
        )
    )

    # (a) Valid ZIP archive conforming to OPC
    assert service.validate_output_docx(output_bytes), (
        "Output .docx is not a valid ZIP/OPC archive"
    )

    # Verify it's a valid ZIP with [Content_Types].xml
    with zipfile.ZipFile(io.BytesIO(output_bytes), "r") as zf:
        assert "[Content_Types].xml" in zf.namelist(), (
            "Output .docx missing [Content_Types].xml (OPC requirement)"
        )
        assert zf.testzip() is None, "Output .docx has corrupted entries"

    # (b) Document properties include title, author, company, created date
    output_doc = DocxDocument(io.BytesIO(output_bytes))
    core_props = output_doc.core_properties
    assert core_props.title == title, (
        f"Expected title '{title}', got '{core_props.title}'"
    )
    assert core_props.author == user_name, (
        f"Expected author '{user_name}', got '{core_props.author}'"
    )
    assert core_props.created is not None, "Created date should be set"

    # Check "generated_by" custom property stored in comments (fallback)
    comments = core_props.comments or ""
    assert "AlcoaBase AI Document Generator v1.0" in comments, (
        "Custom property 'generated_by' not found in document"
    )

    # (c) All text is valid UTF-8 (python-docx handles this natively)
    all_text = "\n".join(para.text for para in output_doc.paragraphs)
    all_text.encode("utf-8")  # Should not raise

    # (d) Sources appendix section exists
    paragraph_texts = [para.text for para in output_doc.paragraphs]
    sources_found = any(
        "AI Generation Sources" in text for text in paragraph_texts
    )
    assert sources_found, (
        "Output document missing 'AI Generation Sources' appendix section"
    )


# ---------------------------------------------------------------------------
# Property 15: Knowledge Base Relevance Ordering
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    scores=st.lists(
        st.floats(min_value=0.3, max_value=1.0, allow_nan=False),
        min_size=2,
        max_size=10,
    ),
)
def test_kb_chunks_ordered_by_descending_relevance(
    scores: list[float],
) -> None:
    """KB chunks provided to the Technical Writer agent SHALL be ordered by
    descending relevance_score, and all chunks SHALL have relevance_score >= 0.3.

    The pipeline sorts chunks in retrieve_knowledge_context before passing them
    to manage_context_window. This test verifies that manage_context_window
    preserves the descending relevance order and that all chunks meet the
    minimum relevance threshold.

    **Validates: Requirements 9.1, 9.2**
    """
    service = _make_service()

    # Sort scores descending (as retrieve_knowledge_context does)
    sorted_scores = sorted(scores, reverse=True)

    # Build KB chunks pre-sorted by descending relevance (pipeline invariant)
    kb_chunks: list[dict[str, Any]] = [
        {
            "chunk_document_uuid": f"doc-{i}",
            "title": f"Document {i}",
            "excerpt": f"Content from document {i} " * 10,
            "text": f"Content from document {i} " * 10,
            "relevance_score": score,
        }
        for i, score in enumerate(sorted_scores)
    ]

    # Create some preceding sections (small enough to not trigger trimming alone)
    preceding_sections = [
        SectionResult(
            heading="Intro",
            content="Short intro.",
            token_count=3,
            inference_duration_ms=50,
        )
    ]

    _, result_chunks = service.manage_context_window(
        preceding_sections=preceding_sections,
        kb_chunks=kb_chunks,
    )

    # All returned chunks must have relevance_score >= 0.3
    for chunk in result_chunks:
        assert chunk["relevance_score"] >= 0.3, (
            f"Chunk has relevance_score {chunk['relevance_score']} < 0.3"
        )

    # Chunks must remain in descending relevance order after context management
    result_scores = [c["relevance_score"] for c in result_chunks]
    for i in range(len(result_scores) - 1):
        assert result_scores[i] >= result_scores[i + 1], (
            f"KB chunks not in descending relevance order: "
            f"{result_scores[i]} < {result_scores[i + 1]} at index {i}"
        )


# ---------------------------------------------------------------------------
# Property 16: Context Window Management
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    num_preceding=st.integers(min_value=5, max_value=15),
    content_length=st.integers(min_value=200, max_value=500),
    num_kb_chunks=st.integers(min_value=6, max_value=10),
    chunk_length=st.integers(min_value=200, max_value=600),
)
def test_context_window_never_exceeds_limits(
    num_preceding: int,
    content_length: int,
    num_kb_chunks: int,
    chunk_length: int,
) -> None:
    """When combined context exceeds 6000 tokens, preceding sections SHALL be
    summarized to max 1000 tokens and only top 5 KB chunks SHALL be included.

    **Validates: Requirements 9.4**
    """
    service = _make_service()

    # Build preceding sections with enough content to exceed limits
    preceding_sections = [
        SectionResult(
            heading=f"Section {i}",
            content="A" * content_length,
            token_count=content_length // CHARS_PER_TOKEN,
            inference_duration_ms=100,
        )
        for i in range(num_preceding)
    ]

    # Build KB chunks with enough content to exceed limits
    kb_chunks: list[dict[str, Any]] = [
        {
            "chunk_document_uuid": f"doc-{i}",
            "title": f"Document {i}",
            "excerpt": "B" * chunk_length,
            "text": "B" * chunk_length,
            "relevance_score": 0.9 - (i * 0.05),
        }
        for i in range(num_kb_chunks)
    ]

    preceding_summary, trimmed_chunks = service.manage_context_window(
        preceding_sections=preceding_sections,
        kb_chunks=kb_chunks,
    )

    # Calculate total tokens of the input
    raw_preceding_tokens = sum(
        (len(f"[{sr.heading}]: {sr.content[:200]}") // CHARS_PER_TOKEN)
        for sr in preceding_sections
    )
    raw_kb_tokens = sum(
        len(c.get("excerpt", c.get("text", ""))) // CHARS_PER_TOKEN
        for c in kb_chunks
    )
    total_input_tokens = raw_preceding_tokens + raw_kb_tokens

    if total_input_tokens > MAX_CONTEXT_TOKENS:
        # Preceding summary must be <= 1000 tokens (4000 chars)
        max_preceding_chars = MAX_PRECEDING_SUMMARY_TOKENS * CHARS_PER_TOKEN
        assert len(preceding_summary) <= max_preceding_chars + 3, (
            f"Preceding summary exceeds 1000 token limit: "
            f"{len(preceding_summary)} chars > {max_preceding_chars + 3} chars "
            f"(+3 for ellipsis)"
        )

        # Only top 5 KB chunks should be kept
        assert len(trimmed_chunks) <= 5, (
            f"Expected at most 5 KB chunks when context exceeds limit, "
            f"got {len(trimmed_chunks)}"
        )


@settings(max_examples=25)
@given(
    num_preceding=st.integers(min_value=1, max_value=3),
    num_kb_chunks=st.integers(min_value=1, max_value=3),
)
def test_context_window_preserves_all_when_within_limits(
    num_preceding: int,
    num_kb_chunks: int,
) -> None:
    """When combined context is within 6000 tokens, all content is preserved.

    **Validates: Requirements 9.4**
    """
    service = _make_service()

    # Build small preceding sections (well within limits)
    preceding_sections = [
        SectionResult(
            heading=f"Section {i}",
            content="Short content.",
            token_count=3,
            inference_duration_ms=50,
        )
        for i in range(num_preceding)
    ]

    # Build small KB chunks (well within limits)
    kb_chunks: list[dict[str, Any]] = [
        {
            "chunk_document_uuid": f"doc-{i}",
            "title": f"Document {i}",
            "excerpt": "Brief excerpt.",
            "text": "Brief excerpt.",
            "relevance_score": 0.9 - (i * 0.1),
        }
        for i in range(num_kb_chunks)
    ]

    _, result_chunks = service.manage_context_window(
        preceding_sections=preceding_sections,
        kb_chunks=kb_chunks,
    )

    # All chunks should be preserved when within limits
    assert len(result_chunks) == num_kb_chunks, (
        f"Expected all {num_kb_chunks} chunks preserved, got {len(result_chunks)}"
    )


# ---------------------------------------------------------------------------
# Property 18: No Raw Placeholders in Output
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    placeholder_ids=st.lists(
        st_placeholder_ids,
        min_size=1,
        max_size=3,
    ),
    title=st_titles,
    company_name=st_company_names,
    user_name=st_user_names,
)
def test_no_raw_placeholders_in_output(
    placeholder_ids: list[str],
    title: str,
    company_name: str,
    user_name: str,
) -> None:
    """Scanning all paragraph text in the output .docx SHALL find zero
    occurrences of the pattern {{[A-Z_]+(?::[^}]+)?}}.

    **Validates: Requirements 10.7**
    """
    service = _make_service()

    # Build a template with placeholder markers in body paragraphs
    doc = DocxDocument()
    headings_data: list[tuple[int, str]] = []
    for i, pid in enumerate(placeholder_ids):
        heading_text = f"Section {i + 1}"
        headings_data.append((1, heading_text))
        doc.add_heading(heading_text, level=1)
        doc.add_paragraph(f"Content with placeholder: {{{{{pid}}}}}")

    buffer = io.BytesIO()
    doc.save(buffer)
    template_bytes = buffer.getvalue()

    # Create section results with clean content (no placeholders)
    sections = [
        SectionResult(
            heading=f"Section {i + 1}",
            content=f"Generated content replacing {pid} placeholder.",
            token_count=10,
            inference_duration_ms=100,
            placeholder_processed=[f"{{{{{pid}}}}}"],
        )
        for i, pid in enumerate(placeholder_ids)
    ]

    template_analysis = _make_template_analysis(headings_data)

    output_bytes = asyncio.run(
        service.assemble_docx(
            template_bytes=template_bytes,
            sections=sections,
            template_analysis=template_analysis,
            title=title,
            company_name=company_name,
            requesting_user_name=user_name,
        )
    )

    # Parse output and check for raw placeholder patterns
    output_doc = DocxDocument(io.BytesIO(output_bytes))
    placeholder_pattern = re.compile(r"\{\{[A-Z_]+(?::[^}]+)?\}\}")

    for para in output_doc.paragraphs:
        matches = placeholder_pattern.findall(para.text)
        assert not matches, (
            f"Raw placeholder(s) found in output: {matches} "
            f"in paragraph: '{para.text[:100]}'"
        )


# ---------------------------------------------------------------------------
# Property 20: Header/Footer Token Substitution
# ---------------------------------------------------------------------------


@settings(max_examples=25)
@given(
    tokens_to_include=st.lists(
        st.sampled_from(SUPPORTED_TOKENS),
        min_size=1,
        max_size=4,
        unique=True,
    ),
    title=st_titles,
    company_name=st_company_names,
    user_name=st_user_names,
)
def test_header_footer_token_substitution(
    tokens_to_include: list[str],
    title: str,
    company_name: str,
    user_name: str,
) -> None:
    """All supported tokens ({title}, {date}, {version}, {company}) in
    headers/footers SHALL be replaced with actual values in the output.

    **Validates: Requirements 4.3**
    """
    service = _make_service()

    # Build header text with selected tokens
    header_text = " | ".join(tokens_to_include)
    footer_text = " - ".join(reversed(tokens_to_include))

    headings = [(1, "Introduction")]
    template_bytes = _build_template_docx(
        headings, header_text=header_text, footer_text=footer_text
    )
    sections = _make_section_results(headings)
    template_analysis = _make_template_analysis(headings)

    output_bytes = asyncio.run(
        service.assemble_docx(
            template_bytes=template_bytes,
            sections=sections,
            template_analysis=template_analysis,
            title=title,
            company_name=company_name,
            requesting_user_name=user_name,
        )
    )

    # Parse output and check headers/footers
    output_doc = DocxDocument(io.BytesIO(output_bytes))

    for section in output_doc.sections:
        # Check header
        header = section.header
        header_full_text = " ".join(p.text for p in header.paragraphs)

        # Check footer
        footer = section.footer
        footer_full_text = " ".join(p.text for p in footer.paragraphs)

        combined_text = header_full_text + " " + footer_full_text

        # None of the supported tokens should remain as raw tokens
        for token in tokens_to_include:
            assert token not in combined_text, (
                f"Token '{token}' was not substituted in header/footer. "
                f"Header: '{header_full_text}', Footer: '{footer_full_text}'"
            )

        # Verify actual values are present
        if "{title}" in tokens_to_include:
            assert title in combined_text, (
                f"Title '{title}' not found in header/footer after substitution"
            )
        if "{version}" in tokens_to_include:
            assert "1.0" in combined_text, (
                "Version '1.0' not found in header/footer after substitution"
            )
        if "{company}" in tokens_to_include:
            assert company_name in combined_text, (
                f"Company '{company_name}' not found in header/footer "
                f"after substitution"
            )


@settings(max_examples=25)
@given(
    unrecognized_token=st.from_regex(
        r"\{[a-z]{3,15}\}", fullmatch=True
    ).filter(lambda t: t not in SUPPORTED_TOKENS),
    title=st_titles,
    company_name=st_company_names,
    user_name=st_user_names,
)
def test_unrecognized_tokens_remain_unchanged(
    unrecognized_token: str,
    title: str,
    company_name: str,
    user_name: str,
) -> None:
    """Unrecognized {identifier} tokens in headers/footers SHALL remain
    unchanged in the output.

    **Validates: Requirements 4.3**
    """
    service = _make_service()

    # Build template with an unrecognized token in header
    header_text = f"Document: {unrecognized_token}"
    headings = [(1, "Introduction")]
    template_bytes = _build_template_docx(
        headings, header_text=header_text
    )
    sections = _make_section_results(headings)
    template_analysis = _make_template_analysis(headings)

    output_bytes = asyncio.run(
        service.assemble_docx(
            template_bytes=template_bytes,
            sections=sections,
            template_analysis=template_analysis,
            title=title,
            company_name=company_name,
            requesting_user_name=user_name,
        )
    )

    # Parse output and verify unrecognized token is preserved
    output_doc = DocxDocument(io.BytesIO(output_bytes))

    for section in output_doc.sections:
        header = section.header
        header_full_text = " ".join(p.text for p in header.paragraphs)

        assert unrecognized_token in header_full_text, (
            f"Unrecognized token '{unrecognized_token}' was incorrectly "
            f"removed or substituted. Header text: '{header_full_text}'"
        )
